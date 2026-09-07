"""
title: Office Export Buttons
author: aiops
version: 1.0.0
"""

import base64
import json
from datetime import datetime

import aiohttp
from pydantic import BaseModel, Field


def _icon(letter: str) -> str:
    svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24">'
        '<rect x="3" y="2" width="18" height="20" rx="3" fill="black"/>'
        f'<text x="12" y="17" text-anchor="middle" font-family="Arial" '
        f'font-size="13" font-weight="700" fill="white">{letter}</text></svg>'
    )
    return "data:image/svg+xml," + svg.replace("#", "%23").replace(" ", "%20")


actions = [
    {"id": "word", "name": "生成 Word", "icon_url": _icon("W")},
    {"id": "excel", "name": "生成 Excel", "icon_url": _icon("X")},
    {"id": "powerpoint", "name": "生成 PPT", "icon_url": _icon("P")},
]


GENERATOR = r'''
import base64
import json
import os
import re
from datetime import datetime

payload = json.loads(base64.b64decode("__PAYLOAD__").decode("utf-8"))
kind = payload["kind"]
content = payload["content"].replace("\r\n", "\n").strip()
output_path = payload["output_path"]
os.makedirs(os.path.dirname(output_path), exist_ok=True)


def clean_inline(value):
    value = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"(`{1,3}|\*\*|__|~~)", "", value)
    return value.strip()


def split_cells(line):
    return [clean_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]


def is_separator(line):
    cells = split_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def markdown_tables(lines):
    tables = []
    i = 0
    while i + 1 < len(lines):
        if "|" in lines[i] and is_separator(lines[i + 1]):
            rows = [split_cells(lines[i])]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_cells(lines[i]))
                i += 1
            tables.append(rows)
        else:
            i += 1
    return tables


def document_title():
    for line in content.splitlines():
        text = clean_inline(re.sub(r"^#{1,6}\s+", "", line).strip())
        if text:
            return text[:80]
    return "AI 内容导出"


def set_docx_font(run, name="Microsoft YaHei", size=None, bold=None):
    from docx.oxml.ns import qn
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        from docx.shared import Pt
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def generate_docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = __import__("docx.shared", fromlist=["Pt"]).Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_docx_font(title.add_run(document_title()), size=20, bold=True)
    stamp = doc.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_docx_font(stamp.add_run(datetime.now().strftime("生成时间：%Y-%m-%d %H:%M")), size=9)

    lines = content.splitlines()
    i = 0
    in_code = False
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if i + 1 < len(lines) and "|" in raw and is_separator(lines[i + 1]):
            rows = [split_cells(raw)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_cells(lines[i]))
                i += 1
            width = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    run = table.cell(r, c).paragraphs[0].add_run(value)
                    set_docx_font(run, size=9.5, bold=(r == 0))
            continue
        if not stripped:
            doc.add_paragraph()
        elif in_code:
            p = doc.add_paragraph()
            set_docx_font(p.add_run(raw), name="Consolas", size=9)
        else:
            heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
            bullet = re.match(r"^\s*[-*+]\s+(.+)$", raw)
            numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", raw)
            if heading:
                level = min(len(heading.group(1)), 3)
                p = doc.add_heading(clean_inline(heading.group(2)), level=level)
                for run in p.runs:
                    set_docx_font(run, bold=True)
            elif bullet:
                p = doc.add_paragraph(style="List Bullet")
                set_docx_font(p.add_run(clean_inline(bullet.group(1))), size=10.5)
            elif numbered:
                p = doc.add_paragraph(style="List Number")
                set_docx_font(p.add_run(clean_inline(numbered.group(1))), size=10.5)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0.74)
                set_docx_font(p.add_run(clean_inline(stripped)), size=10.5)
        i += 1
    doc.save(output_path)


def generate_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "内容"
    ws.sheet_view.showGridLines = False
    ws["A1"] = document_title()
    ws["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A2"] = datetime.now().strftime("生成时间：%Y-%m-%d %H:%M")
    ws["A2"].font = Font(name="Microsoft YaHei", size=9, color="666666")
    ws.column_dimensions["A"].width = 100

    row = 4
    for line in content.splitlines():
        if not is_separator(line):
            value = clean_inline(re.sub(r"^#{1,6}\s+", "", line))
            if value:
                ws.cell(row=row, column=1, value=value)
                ws.cell(row=row, column=1).font = Font(name="Microsoft YaHei", size=10.5)
                ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
                row += 1

    for index, rows in enumerate(markdown_tables(content.splitlines()), 1):
        sheet = wb.create_sheet(f"表{index}")
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for r, values in enumerate(rows, 1):
            for c, value in enumerate(values, 1):
                cell = sheet.cell(r, c, value)
                cell.font = Font(name="Microsoft YaHei", size=10, bold=(r == 1), color="FFFFFF" if r == 1 else "000000")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if r == 1:
                    cell.fill = PatternFill("solid", fgColor="4472C4")
        for c in range(1, max(len(row) for row in rows) + 1):
            width = max(len(str(sheet.cell(r, c).value or "")) for r in range(1, len(rows) + 1))
            sheet.column_dimensions[__import__("openpyxl.utils", fromlist=["get_column_letter"]).get_column_letter(c)].width = min(max(width * 1.5, 12), 45)
    wb.save(output_path)


def ppt_sections():
    sections = []
    current_title = None
    current_lines = []
    for raw in content.splitlines():
        heading = re.match(r"^#{1,3}\s+(.+)$", raw.strip())
        if heading:
            if current_title or current_lines:
                sections.append((current_title or "内容", current_lines))
            current_title = clean_inline(heading.group(1))
            current_lines = []
        elif raw.strip():
            current_lines.append(clean_inline(re.sub(r"^\s*(?:[-*+]|\d+[.)])\s+", "", raw.strip())))
    if current_title or current_lines:
        sections.append((current_title or "内容", current_lines))
    if not sections:
        sections = [("内容", [content])]
    expanded = []
    for title, lines in sections:
        chunks = [lines[i:i + 6] for i in range(0, max(len(lines), 1), 6)] or [[]]
        for number, chunk in enumerate(chunks, 1):
            expanded.append((title if number == 1 else f"{title}（续）", chunk))
    return expanded


def generate_pptx():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor(31, 78, 121)
    cover.shapes.title.text = document_title()
    cover.placeholders[1].text = datetime.now().strftime("AI 内容导出  ·  %Y-%m-%d")
    for shape in cover.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
                run.font.color.rgb = RGBColor(255, 255, 255)

    for title, lines in ppt_sections():
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title[:60]
        body = slide.placeholders[1].text_frame
        body.clear()
        for index, line in enumerate(lines or [""], 0):
            p = body.paragraphs[0] if index == 0 else body.add_paragraph()
            p.text = line[:280]
            p.level = 0
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(20)
            p.space_after = Pt(10)
        for run in slide.shapes.title.text_frame.paragraphs[0].runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(28)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
    prs.save(output_path)


if kind == "word":
    generate_docx()
elif kind == "excel":
    generate_xlsx()
elif kind == "powerpoint":
    generate_pptx()
else:
    raise ValueError(f"Unsupported format: {kind}")

print(json.dumps({"path": output_path, "kind": kind}, ensure_ascii=False))
'''


class Action:
    actions = actions

    class Valves(BaseModel):
        terminal_url: str = Field(
            default="http://127.0.0.1:8000",
            description="Open Terminal URL reachable from the Open WebUI backend.",
        )
        terminal_api_key: str = Field(
            default="",
            description="Open Terminal Bearer API key.",
            json_schema_extra={"input": {"type": "password"}},
        )
        output_dir: str = Field(
            default="/home/user/exports",
            description="Directory where generated Office files are saved.",
        )
        timeout_seconds: int = Field(default=120, description="Maximum generation time in seconds.")
        priority: int = Field(default=20, description="Button display order; lower appears first.")

    def __init__(self):
        self.valves = self.Valves()

    @staticmethod
    def _message_content(body: dict) -> str:
        for message in reversed(body.get("messages", [])):
            if message.get("role") == "assistant" and str(message.get("content", "")).strip():
                return str(message["content"]).strip()
        return ""

    async def _emit(self, emitter, event_type: str, data: dict):
        if emitter:
            await emitter({"type": event_type, "data": data})

    async def action(self, body: dict, __id__=None, __event_emitter__=None, **_):
        labels = {"word": ("Word", "docx"), "excel": ("Excel", "xlsx"), "powerpoint": ("PPT", "pptx")}
        if __id__ not in labels:
            raise ValueError("不支持的 Office 导出格式")
        if not self.valves.terminal_api_key.strip():
            raise ValueError("请先在此 Action 的 Valves 中配置 Open Terminal API Key")

        content = self._message_content(body)
        if not content:
            raise ValueError("当前回答没有可导出的文本内容")

        label, extension = labels[__id__]
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = f"{self.valves.output_dir.rstrip('/')}/AI内容_{stamp}.{extension}"
        payload = base64.b64encode(
            json.dumps({"kind": __id__, "content": content, "output_path": output_path}, ensure_ascii=False).encode("utf-8")
        ).decode("ascii")
        program = GENERATOR.replace("__PAYLOAD__", payload)
        encoded_program = base64.b64encode(program.encode("utf-8")).decode("ascii")
        command = f'python -c "import base64;exec(base64.b64decode(\'{encoded_program}\'))"'

        await self._emit(
            __event_emitter__,
            "status",
            {"description": f"正在生成 {label} 文件…", "done": False},
        )

        base_url = self.valves.terminal_url.rstrip("/")
        headers = {"Authorization": f"Bearer {self.valves.terminal_api_key.strip()}"}
        timeout = aiohttp.ClientTimeout(total=max(self.valves.timeout_seconds + 10, 30))
        try:
            async with aiohttp.ClientSession(timeout=timeout) as session:
                async with session.post(
                    f"{base_url}/execute",
                    params={"wait": self.valves.timeout_seconds},
                    headers=headers,
                    json={"command": command, "cwd": "/home/user"},
                ) as response:
                    result = await response.json(content_type=None)
                    if response.status >= 400:
                        raise RuntimeError(result.get("detail") or result)
                if result.get("status") != "done" or result.get("exit_code") != 0:
                    output = "".join(item.get("data", "") for item in result.get("output", []))
                    raise RuntimeError(output.strip() or f"Open Terminal 返回状态：{result.get('status')}")
                async with session.get(
                    f"{base_url}/files/display",
                    params={"path": output_path},
                    headers=headers,
                ) as response:
                    display_result = await response.json(content_type=None)
                    if response.status >= 400 or display_result.get("exists") is False:
                        raise RuntimeError("文件已生成，但 Open Terminal 无法显示该文件")

            await self._emit(__event_emitter__, "terminal:run_command", {})
            await self._emit(__event_emitter__, "terminal:display_file", {"path": output_path})
            await self._emit(
                __event_emitter__,
                "notification",
                {"type": "success", "content": f"{label} 已生成：{output_path}"},
            )
            await self._emit(
                __event_emitter__,
                "status",
                {"description": f"{label} 文件生成完成", "done": True},
            )
            return None
        except Exception as error:
            await self._emit(
                __event_emitter__,
                "notification",
                {"type": "error", "content": f"{label} 生成失败：{error}"},
            )
            await self._emit(
                __event_emitter__,
                "status",
                {"description": f"{label} 文件生成失败", "done": True, "error": True},
            )
            raise
